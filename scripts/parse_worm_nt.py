"""Parse WormAtlas neurotransmitter Excel and create neuron->NT mapping.

Source: https://www.wormatlas.org/images/Ce_NTtables_Loer&Rand2022.xlsx
Sheet: 'All neurons, sorted by name' (header at row 8, data from row 9)

Output: data/worm/neuron_nt_mapping.json
"""
import json
import os
import sys
from pathlib import Path

import openpyxl

HERE = Path(__file__).resolve().parent
ROOT = HERE.parent
EXCEL_PATH = ROOT / "tmp" / "worm_nt.xlsx"
OUT_PATH = ROOT / "data" / "worm" / "neuron_nt_mapping.json"

# Normalized neurotransmitter names
NT_NORMALIZE = {
    "Acetylcholine (ACh)": "ACh",
    "Glutamate (Glu)": "Glu",
    "GABA": "GABA",
    "Dopamine (DA)": "DA",
    "Serotonin / 5HT": "5HT",
    "Serotonin (5HT)": "5HT",
    "Tyramine (TA)": "TA",
    "Octopamine (OA)": "OA",
    "Tyramine / Octopamine (TA/OA)": "TA/OA",
}


def normalize_nt(name):
    """Normalize neurotransmitter name."""
    if not name:
        return None
    name = name.strip()
    # Direct match
    if name in NT_NORMALIZE:
        return NT_NORMALIZE[name]
    # Partial match
    for key, val in NT_NORMALIZE.items():
        if key.startswith(name) or name.startswith(key.split(" (")[0]):
            return val
    return name


def parse_excel():
    """Parse the Excel file and return neuron->NT mapping."""
    if not EXCEL_PATH.exists():
        print(f"ERROR: Excel file not found at {EXCEL_PATH}", file=sys.stderr)
        print("  Download from: https://www.wormatlas.org/images/Ce_NTtables_Loer&Rand2022.xlsx", file=sys.stderr)
        sys.exit(1)

    wb = openpyxl.load_workbook(EXCEL_PATH, data_only=True)
    ws = wb["All neurons, sorted by name"]

    # Header at row 8, data from row 9
    # Col 4: Neuron name, Col 6: NT1, Col 7: NT2
    mapping = {}
    current_class = None

    for row in ws.iter_rows(min_row=9, max_row=ws.max_row, values_only=True):
        neuron_name = row[3]  # Col D
        nt1 = row[5]          # Col F
        nt2 = row[6]          # Col G

        if not neuron_name:
            continue

        neuron_name = str(neuron_name).strip()
        if not neuron_name:
            continue

        # Normalize neurotransmitters
        nt1_norm = normalize_nt(nt1)
        nt2_norm = normalize_nt(nt2)

        # Collect neurotransmitters
        nts = []
        if nt1_norm:
            nts.append(nt1_norm)
        if nt2_norm and nt2_norm != nt1_norm:
            nts.append(nt2_norm)

        mapping[neuron_name] = {
            "neurotransmitters": nts,
            "nt1": nt1_norm,
            "nt2": nt2_norm,
        }

    return mapping


def summarize(mapping):
    """Summarize neurotransmitter distribution."""
    from collections import Counter

    nt_counter = Counter()
    for neuron, data in mapping.items():
        for nt in data["neurotransmitters"]:
            nt_counter[nt] += 1

    return dict(nt_counter.most_common())


def main():
    print("Parsing WormAtlas neurotransmitter Excel...")
    mapping = parse_excel()
    print(f"  Parsed {len(mapping)} neurons")

    # Summarize
    summary = summarize(mapping)
    print("  Neurotransmitter distribution:")
    for nt, count in summary.items():
        print(f"    {nt:10s}: {count} neurons")

    # Save
    OUT_PATH.parent.mkdir(parents=True, exist_ok=True)
    output = {
        "source": "WormAtlas Loer & Rand 2022",
        "source_url": "https://www.wormatlas.org/images/Ce_NTtables_Loer&Rand2022.xlsx",
        "n_neurons": len(mapping),
        "summary": summary,
        "neurons": mapping,
    }
    OUT_PATH.write_text(json.dumps(output, indent=2), encoding="utf-8")
    print(f"\nSaved to {OUT_PATH}")


if __name__ == "__main__":
    main()
